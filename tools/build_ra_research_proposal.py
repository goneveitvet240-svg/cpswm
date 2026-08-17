#!/usr/bin/env python3
"""Build the RA research proposal DOCX and its architecture figure."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "ra_research_proposal_continual_personalized_world_model.md"
OUT_DIR = ROOT / "output" / "docx"
ASSET_DIR = ROOT / "output" / "assets"
DOCX_PATH = OUT_DIR / "continual_personalized_world_model_RA_research_proposal.docx"
FIGURE_PATH = ASSET_DIR / "cpswm_architecture.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREY = "F4F6F9"
MID_GREY = "D8DEE6"
DARK_GREY = "3B4652"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color=MID_GREY, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_fixed_table_width(table, width_dxa=9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_column_widths(table, widths_in: list[float]) -> None:
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(grid_col)
    for row in table.rows:
        for index, width in enumerate(widths_in):
            if index < len(row.cells):
                row.cells[index].width = Inches(width)
                tc_pr = row.cells[index]._tc.get_or_add_tcPr()
                tc_w = tc_pr.first_child_found_in("w:tcW")
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(int(width * 1440)))
                tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb("6B7280")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(https?://\S+|\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+`)")


def add_inline(paragraph, text: str, *, default_size=11, default_color=DARK_GREY) -> None:
    text = text.replace("\\(", "").replace("\\)", "")
    text = text.replace("\\theta_{perception}", "θ_perception")
    text = text.replace("\\mathcal{G}_{geometry}", "G_geometry")
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            run.font.name = "Calibri"
            run.font.size = Pt(default_size)
            run.font.color.rgb = rgb(default_color)
        token = match.group(0)
        if token.startswith("http"):
            url = token.rstrip(".,;)")
            trailing = token[len(url) :]
            add_hyperlink(paragraph, url, url)
            if trailing:
                run = paragraph.add_run(trailing)
                run.font.name = "Calibri"
                run.font.size = Pt(default_size)
                run.font.color.rgb = rgb(default_color)
        else:
            content = token
            bold = token.startswith("**")
            italic = token.startswith("*") and not bold
            code = token.startswith("`")
            content = token[2:-2] if bold else token[1:-1]
            run = paragraph.add_run(content)
            run.bold = bold
            run.italic = italic
            run.font.name = "Consolas" if code else "Calibri"
            run.font.size = Pt(9.5 if code else default_size)
            run.font.color.rgb = rgb(DARK_BLUE if bold else default_color)
            if code:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), LIGHT_GREY)
                run._r.get_or_add_rPr().append(shading)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.name = "Calibri"
        run.font.size = Pt(default_size)
        run.font.color.rgb = rgb(default_color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(DARK_GREY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = rgb(BLUE)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = rgb(BLUE)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = rgb(DARK_BLUE)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = rgb(DARK_GREY)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    props = doc.core_properties
    props.title = "Continual Personalized Semantic World Models for Long-Term Household Robots"
    props.subject = "Research Proposal for a Prospective Remote Research Assistantship"
    props.author = "Pang Wei"
    props.keywords = "embodied AI, long-term robot memory, semantic world models, continual learning"


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("PANG WEI  |  CPSWM RESEARCH PROPOSAL")
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = rgb("6B7280")
    p.paragraph_format.space_after = Pt(0)

    footer = section.footer
    add_page_field(footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CONTINUAL PERSONALIZED\nSEMANTIC WORLD MODELS")
    run.font.name = "Calibri"
    run.font.size = Pt(27)
    run.font.bold = True
    run.font.color.rgb = rgb(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("for Long-Term Household Robots")
    run.font.name = "Calibri Light"
    run.font.size = Pt(18)
    run.font.color.rgb = rgb(DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("Research Proposal for a Prospective Remote Research Assistantship")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = rgb("5B6571")

    rule = doc.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    rule.autofit = False
    rule.columns[0].width = Inches(4.6)
    cell = rule.cell(0, 0)
    set_cell_shading(cell, BLUE)
    set_cell_margins(cell, top=12, bottom=12, start=0, end=0)
    cell.text = ""
    set_table_borders(rule, BLUE, "0")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    metadata = [
        ("Applicant", "Pang Wei"),
        ("Institution", "Qingdao University"),
        ("Programme", "B.Eng. in Intelligent Manufacturing Engineering (expected June 2027)"),
        ("Research areas", "Embodied AI · Continual Learning · Long-Term Robot Memory"),
        ("Email", "m15964096907@outlook.com"),
        ("Date", "August 2026"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_width(table, 7920)
    set_column_widths(table, [1.35, 4.15])
    set_table_borders(table, WHITE, "0")
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(row.cells[0], 80, 60, 80, 100)
        set_cell_margins(row.cells[1], 80, 100, 80, 60)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r0 = p0.add_run(label.upper())
        r0.font.name = "Calibri"
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = rgb(BLUE)
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(value)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = rgb(DARK_GREY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Persistent instance identity  ·  Person-conditioned routines  ·  "
        "Evidence-grounded belief revision  ·  Memory–action writeback"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = rgb("6B7280")

    doc.add_page_break()


def make_architecture_figure(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    scale = 1
    image = Image.new("RGB", (2080, 1240), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    if not Path(regular_path).exists():
        regular_path = "/System/Library/Fonts/Helvetica.ttc"
        bold_path = regular_path
    title_font = ImageFont.truetype(bold_path, 30)
    text_font = ImageFont.truetype(regular_path, 24)
    label_font = ImageFont.truetype(regular_path, 20)

    def hx(value: str):
        return f"#{value}"

    def box(x, y, w, h, title, subtitle, fill, edge=BLUE, title_color=WHITE, text_color=DARK_GREY):
        xy = (x * scale, y * scale, (x + w) * scale, (y + h) * scale)
        draw.rounded_rectangle(xy, radius=22, fill=hx(fill), outline=hx(edge), width=3)
        draw.text(((x + 34) * scale, (y + 22) * scale), title, font=title_font, fill=hx(title_color))
        draw.multiline_text(((x + 34) * scale, (y + 70) * scale), subtitle, font=text_font, fill=hx(text_color), spacing=8)

    box(110, 70, 1860, 145, "A  CONTRACTS & RUNTIME  ·  M01–M04", "Ontology · identity/time/coordinates · persistence · orchestration", BLUE, text_color=WHITE)
    box(110, 265, 1860, 165, "B  EMBODIED OBSERVATION & MAPPING  ·  M05–M12", "Sensors → SLAM / 3D structure → observability → semantics → instance & interaction candidates", LIGHT_BLUE, title_color=DARK_BLUE)
    box(110, 490, 1860, 205, "C  LONG-TERM WORLD MODEL & REASONING  ·  M13–M19", "Canonical assertions + episodic events + provenance\n→ derived multi-hypothesis belief → habits → hidden events → memory lifecycle", DARK_BLUE, edge=DARK_BLUE, text_color=WHITE)
    box(110, 760, 870, 175, "D  QUERY & EXPLANATION  ·  M20–M22", "Typed retrieval · language grounding\nevidence-based explanation & correction", LIGHT_GREY, title_color=DARK_BLUE)
    box(1100, 760, 870, 175, "E  PLANNING & ACTION  ·  M23–M27", "Active sensing · navigation · manipulation\nclosed-loop evidence writeback", LIGHT_GREY, title_color=DARK_BLUE)
    box(110, 1010, 1860, 145, "F  SIMULATION, BENCHMARK & GOVERNANCE  ·  M28–M32", "Privacy · symbolic / embodied simulation · synthetic routines · benchmark · evaluation", "E7EDF3", title_color=DARK_BLUE)

    def arrow(start, end, label=None, label_xy=None):
        sx, sy = start[0] * scale, start[1] * scale
        ex, ey = end[0] * scale, end[1] * scale
        draw.line((sx, sy, ex, ey), fill="#6B7C8F", width=5)
        import math
        angle = math.atan2(ey - sy, ex - sx)
        length = 22
        spread = 0.55
        p1 = (ex - length * math.cos(angle - spread), ey - length * math.sin(angle - spread))
        p2 = (ex - length * math.cos(angle + spread), ey - length * math.sin(angle + spread))
        draw.polygon([(ex, ey), p1, p2], fill="#6B7C8F")
        if label and label_xy:
            draw.text((label_xy[0] * scale, label_xy[1] * scale), label, font=label_font, fill="#6B7280", anchor="mm")

    arrow((1040, 215), (1040, 265))
    arrow((1040, 430), (1040, 490), "structured evidence", (1040, 458))
    arrow((620, 695), (545, 760))
    arrow((1460, 695), (1535, 760))
    arrow((980, 848), (1100, 848), "candidate set", (1040, 825))
    arrow((1535, 760), (1440, 695), "action outcomes", (1630, 724))
    arrow((1040, 1010), (1040, 935))
    image.resize((2080, 1240), Image.Resampling.LANCZOS).save(path, format="PNG", optimize=True)


def add_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(set(cell) <= set("-: ") for cell in rows[1]):
        rows.pop(1)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_width(table)
    if cols == 3:
        widths = [1.28, 1.20, 4.02]
    elif cols == 4:
        widths = [1.05, 2.0, 1.55, 1.9]
    else:
        widths = [6.5 / cols] * cols
    set_column_widths(table, widths)
    set_table_borders(table)
    for r_idx, (word_row, src_row) in enumerate(zip(table.rows, rows)):
        prevent_row_split(word_row)
        if r_idx == 0:
            set_repeat_table_header(word_row)
        for c_idx, cell in enumerate(word_row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, LIGHT_GREY if r_idx == 0 else WHITE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            text = src_row[c_idx] if c_idx < len(src_row) else ""
            add_inline(p, text, default_size=9.0, default_color=DARK_GREY)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = rgb(DARK_BLUE)
        if r_idx % 2 == 0 and r_idx > 0:
            for cell in word_row.cells:
                set_cell_shading(cell, "FAFBFC")
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(1)


def add_quote(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_width(table, 8880)
    set_column_widths(table, [6.17])
    set_table_borders(table, LIGHT_BLUE, "0")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11.5)
    r.font.italic = True
    r.font.color.rgb = rgb(DARK_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_architecture_figure(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(FIGURE_PATH), width=Inches(5.75))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run("Figure 1. Six-layer CPSWM architecture and the closed evidence–action loop.")
    r.font.name = "Calibri"
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = rgb("66717D")


def parse_markdown_into_doc(doc: Document, source: str) -> None:
    lines = source.splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "## Abstract")
    i = start
    architecture_figure_added = False
    in_references = False

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            i += 1
            continue

        if stripped.startswith("## "):
            title = stripped[3:]
            p = doc.add_paragraph(title, style="Heading 1")
            in_references = title == "References"
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:], style="Heading 2")
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            if not architecture_figure_added and any("Contracts and runtime foundation" in row for row in table_lines):
                add_architecture_figure(doc)
                architecture_figure_added = True
            continue

        if stripped.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_quote(doc, " ".join(quote_lines))
            continue

        if stripped == "\\[":
            math_lines = []
            i += 1
            while i < len(lines) and lines[i].strip() != "\\]":
                math_lines.append(lines[i].strip())
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            text = " ".join(math_lines)
            text = text.replace("\\mid", "|").replace("\\sum_v", "Σᵥ").replace("\\theta", "θ")
            text = text.replace("\\mathcal{G}", "𝒢").replace("\\,", " ")
            text = re.sub(r"\\text\{([^}]+)\}", r"\1", text)
            text = text.replace("_{perception}", "ₚₑᵣcₑₚₜᵢₒₙ").replace("_{geometry}", "gₑₒₘₑₜᵣy")
            text = text.replace("_{t+1}", "ₜ₊₁").replace("_t", "ₜ")
            r = p.add_run(text)
            r.font.name = "Cambria Math"
            r.font.size = Pt(11.5)
            r.font.color.rgb = rgb(DARK_BLUE)
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet or numbered:
            if bullet:
                content = bullet.group(1)
                p = doc.add_paragraph(style="List Bullet")
            else:
                content = numbered.group(2)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.375)
                p.paragraph_format.first_line_indent = Inches(-0.194)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.208
                nr = p.add_run(f"{numbered.group(1)}. ")
                nr.font.name = "Calibri"
                nr.font.size = Pt(11)
                nr.font.color.rgb = rgb(DARK_GREY)
            add_inline(p, content)
            i += 1
            continue

        para_lines = [stripped]
        hard_break = line.endswith("  ")
        i += 1
        while i < len(lines) and lines[i].strip():
            nxt = lines[i].strip()
            if hard_break or nxt.startswith(("## ", "### ", "|", "> ", "- ", "\\[")) or re.match(r"^\d+\.\s+", nxt):
                break
            para_lines.append(nxt)
            hard_break = lines[i].rstrip().endswith("  ")
            i += 1

        text = " ".join(part.rstrip() for part in para_lines)
        p = doc.add_paragraph()
        if in_references:
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, text, default_size=9.0)
        else:
            p.paragraph_format.widow_control = True
            add_inline(p, text)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_architecture_figure(FIGURE_PATH)

    doc = Document()
    configure_document(doc)
    configure_header_footer(doc)
    add_cover(doc)
    parse_markdown_into_doc(doc, SOURCE.read_text(encoding="utf-8"))

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
