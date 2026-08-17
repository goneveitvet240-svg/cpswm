#!/usr/bin/env python3
"""Build the four-page RA research proposal."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_ra_research_proposal import (
    BLUE,
    DARK_BLUE,
    DARK_GREY,
    configure_document,
    configure_header_footer,
    parse_markdown_into_doc,
    rgb,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "ra_research_proposal_continual_personalized_world_model_4page.md"
OUTPUT = ROOT / "output" / "docx" / "continual_personalized_world_model_RA_research_proposal_4page.docx"


def add_compact_title(doc: Document) -> None:
    # Named override: compact_four_page_title. It keeps the proposal-centerpiece
    # hierarchy without spending a full page on a cover.
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(2)
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Continual Personalized Semantic World Models\nfor Long-Term Household Robots")
    run.font.name = "Calibri"
    run.font.size = Pt(17.5)
    run.font.bold = True
    run.font.color.rgb = rgb(BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(3)
    run = subtitle.add_run("Research Proposal for a Prospective Remote Research Assistantship")
    run.font.name = "Calibri"
    run.font.size = Pt(9.8)
    run.font.italic = True
    run.font.color.rgb = rgb(DARK_BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(8)
    run = meta.add_run(
        "Pang Wei  |  Qingdao University  |  Intelligent Manufacturing Engineering\n"
        "m15964096907@outlook.com  |  August 2026"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(8.8)
    run.font.color.rgb = rgb("66717D")


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    configure_header_footer(doc)
    section = doc.sections[0]
    section.different_first_page_header_footer = False

    # Named override: four_page_application_density. The 1-inch page geometry
    # and grant-proposal visual system remain unchanged; type rhythm is tightened
    # to meet the explicit four-page application limit without deleting scope.
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.10
    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(14.5)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(5)
    h2 = doc.styles["Heading 2"]
    h2.font.size = Pt(11.5)
    h2.paragraph_format.space_before = Pt(7)
    h2.paragraph_format.space_after = Pt(3)
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.size = Pt(10.3)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.08

    # Compact running furniture for a short application document.
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("PANG WEI  |  CPSWM RESEARCH PROPOSAL")
    run.font.name = "Calibri"
    run.font.size = Pt(8.2)
    run.font.bold = True
    run.font.color.rgb = rgb("6B7280")

    add_compact_title(doc)
    parse_markdown_into_doc(doc, SOURCE.read_text(encoding="utf-8"))

    doc.core_properties.title = "Continual Personalized Semantic World Models for Long-Term Household Robots - Four-Page RA Research Proposal"
    doc.core_properties.subject = "Four-page research proposal for a prospective remote research assistantship"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
